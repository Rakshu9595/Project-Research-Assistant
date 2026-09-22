
import os
from docx import Document


def load_docx(file_path):
    """
    Load a DOCX file and extract its elements.
    """

    try:
        # Check whether file exists
        if not os.path.exists(file_path):
            raise FileNotFoundError(
                f"File not found: {file_path}"
            )

        # Open DOCX document
        document = Document(file_path)

        elements = []

        
        # EXTRACT PARAGRAPHS
       

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if not text:
                continue

            # Get paragraph style
            style_name = paragraph.style.name

            # Identify heading
            if style_name.startswith("Heading"):
                element_type = "heading"
            else:
                element_type = "text"

            elements.append({
                "element_type": element_type,
                "content": text,
                "style": style_name
            })

       
        # EXTRACT TABLES
       

        for table_index, table in enumerate(
            document.tables,
            start=1
        ):

            table_data = []

            for row in table.rows:

                row_data = []

                for cell in row.cells:

                    cell_text = cell.text.strip()

                    row_data.append(cell_text)

                table_data.append(row_data)

            if table_data:

                elements.append({
                    "element_type": "table",
                    "table_number": table_index,
                    "content": table_data
                })

   
        # EXTRACT IMAGES
     

        image_count = 0

        for relationship in document.part.rels.values():

            if "image" in relationship.reltype:

                image_count += 1

                elements.append({
                    "element_type": "image",
                    "image_number": image_count,
                    "content": relationship.target_ref
                })

        # RETURN DOCUMENT
     

        return {
            "file_type": "docx",
            "file_path": file_path,
            "elements": elements
        }

    except Exception as error:

        print(
            "❌ DOCX Loading Error:",
            error
        )

        raise Exception(
            f"Failed to load DOCX: {str(error)}"
        )

